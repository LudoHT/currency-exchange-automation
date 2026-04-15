from openpyxl import load_workbook

def write_summary_sheet(template_path, output_path, fx_data: dict):
    wb = load_workbook(template_path)
    ws = wb["Summary"]

    start_row = 41

    currencies_in_order = [
        "AED", "AUD", "CAD", "CHF", "DKK", "EUR", "GBP", "HUF",
        "JPY", "NOK", "NZD", "RUB", "SEK", "SGD", "ZAR",
        "INR", "KRW", "MYR", "HKD"
    ]

    # Left table:
    # B = Currency
    # C = Period Average
    # D = Currency again
    # E = Spot Rate/Date
    for idx, currency in enumerate(currencies_in_order, start=start_row):
        if currency not in fx_data:
            continue
        ws[f"B{idx}"] = currency
        ws[f"C{idx}"] = fx_data[currency]["period_average"]
        ws[f"D{idx}"] = currency
        ws[f"E{idx}"] = fx_data[currency]["spot_rate"]

        ws[f"C{idx}"].number_format = "0.00000"
        ws[f"E{idx}"].number_format = "0.00000"

    # Right table:
    # G = Currency
    # H = Salesforce Update value (spot rate only)
    salesforce_start_row = 41
    salesforce_currencies = [
        "AUD", "CAD", "CHF", "DKK", "EUR", "GBP", "HKD", "INR",
        "JPY", "KRW", "MYR", "NOK", "NZD", "SEK", "SGD"
    ]

    for idx, currency in enumerate(salesforce_currencies, start=salesforce_start_row):
        if currency not in fx_data:
            continue
        ws[f"G{idx}"] = currency
        ws[f"H{idx}"] = fx_data[currency]["spot_rate"]
        ws[f"H{idx}"].number_format = "0.00000"

    wb.save(output_path)
